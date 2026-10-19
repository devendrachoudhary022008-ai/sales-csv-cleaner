
import csv
from pathlib import Path
from datetime import date, datetime
from collections import defaultdict
from openpyxl import load_workbook

# Project folders
BASE_DIR = Path(__file__).resolve().parent
DATA_DIR = BASE_DIR / "data"
OUTPUT_DIR = BASE_DIR / "output"

CSV_FILE = DATA_DIR / "sales.csv"
EXCEL_FILE = DATA_DIR / "sales.xlsx"

REQUIRED_COLUMNS = [
    "order_id",
    "date",
    "product",
    "quantity",
    "unit_price",
]

CLEAN_COLUMNS = REQUIRED_COLUMNS + ["total"]


def clean_text(value):
    """Normalize text, dates, and spreadsheet cell values."""
    if value is None:
        return ""

    if isinstance(value, datetime):
        return value.date().isoformat()

    if isinstance(value, date):
        return value.isoformat()

    return str(value).strip()


def read_input():
    """Read CSV if available; otherwise read Excel."""
    if CSV_FILE.exists():
        print(f"Input format: CSV ({CSV_FILE.name})")

        with CSV_FILE.open(
            "r", newline="", encoding="utf-8-sig"
        ) as file:
            reader = csv.DictReader(file)

            if not reader.fieldnames:
                raise ValueError("CSV file is empty or has no header.")

            headers = [clean_text(name) for name in reader.fieldnames]
            rows = [
                {
                    headers[index]: value
                    for index, value in enumerate(row.values())
                    if index < len(headers) and headers[index]
                }
                for row in reader
            ]

        return headers, rows

    if EXCEL_FILE.exists():
        print(f"Input format: Excel ({EXCEL_FILE.name})")

        workbook = load_workbook(
            EXCEL_FILE, read_only=True, data_only=True
        )

        try:
            sheet = workbook.active
            iterator = sheet.iter_rows(values_only=True)

            first_row = next(iterator, None)
            if not first_row:
                raise ValueError("Excel file is empty or has no header.")

            headers = [clean_text(value) for value in first_row]

            rows = []
            for values in iterator:
                row = {
                    headers[index]: values[index]
                    for index in range(min(len(headers), len(values)))
                    if headers[index]
                }
                rows.append(row)

            return headers, rows
        finally:
            workbook.close()

    raise FileNotFoundError(
        "Input file not found. Add data/sales.csv "
        "or data/sales.xlsx."
    )


def process_sales():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    try:
        headers, input_rows = read_input()
    except (OSError, ValueError, KeyError) as error:
        print(f"ERROR: {error}")
        return
    except Exception as error:
        print(f"ERROR: Could not read input file: {error}")
        return

    missing_columns = [
        column for column in REQUIRED_COLUMNS
        if column not in headers
    ]

    if missing_columns:
        print("ERROR: Missing columns:", ", ".join(missing_columns))
        print("Required columns:", ", ".join(REQUIRED_COLUMNS))
        return

    cleaned_rows = []
    rejected_rows = []
    seen_records = set()
    daily_sales = defaultdict(float)

    for line_number, row in enumerate(input_rows, start=2):
        if not any(clean_text(value) for value in row.values()):
            continue

        original_row = {
            key: clean_text(value)
            for key, value in row.items()
        }

        try:
            order_id = clean_text(row.get("order_id"))
            sale_date = clean_text(row.get("date"))
            product = clean_text(row.get("product"))
            quantity_text = clean_text(row.get("quantity"))
            price_text = clean_text(row.get("unit_price"))

            if not all([
                order_id, sale_date, product,
                quantity_text, price_text,
            ]):
                raise ValueError("Missing required value")

            parsed_date = datetime.strptime(
                sale_date, "%Y-%m-%d"
            ).date()

            # Accept whole-number quantities, including Excel cells
            quantity_number = float(quantity_text)
            if not quantity_number.is_integer():
                raise ValueError("Quantity must be a whole number")

            quantity = int(quantity_number)
            unit_price = float(price_text)

            if quantity <= 0:
                raise ValueError("Quantity must be greater than 0")

            if unit_price < 0:
                raise ValueError("Unit price cannot be negative")

            if not (float("-inf") < unit_price < float("inf")):
                raise ValueError("Unit price must be a finite number")

            record_key = (
                order_id, sale_date, product, quantity, unit_price
            )

            if record_key in seen_records:
                raise ValueError("Duplicate sales record")

            seen_records.add(record_key)

            total = quantity * unit_price
            day = parsed_date.isoformat()

            cleaned_rows.append({
                "order_id": order_id,
                "date": day,
                "product": product,
                "quantity": quantity,
                "unit_price": f"{unit_price:.2f}",
                "total": f"{total:.2f}",
            })

            daily_sales[day] += total

        except (ValueError, TypeError, OverflowError) as error:
            rejected_rows.append({
                **original_row,
                "line_number": line_number,
                "rejection_reason": str(error),
            })

    clean_file = OUTPUT_DIR / "cleaned_sales.csv"
    with clean_file.open("w", newline="", encoding="utf-8") as file:
        writer = csv.DictWriter(file, fieldnames=CLEAN_COLUMNS)
        writer.writeheader()
        writer.writerows(cleaned_rows)

    rejected_file = OUTPUT_DIR / "rejected_rows.csv"
    rejected_columns = REQUIRED_COLUMNS + [
        "line_number", "rejection_reason"
    ]

    with rejected_file.open("w", newline="", encoding="utf-8") as file:
        writer = csv.DictWriter(
            file,
            fieldnames=rejected_columns,
            extrasaction="ignore",
        )
        writer.writeheader()
        writer.writerows(rejected_rows)

    report_file = OUTPUT_DIR / "daily_sales_report.csv"
    with report_file.open("w", newline="", encoding="utf-8") as file:
        writer = csv.DictWriter(
            file,
            fieldnames=["date", "total_sales", "valid_orders"],
        )
        writer.writeheader()

        for day in sorted(daily_sales):
            order_count = sum(
                1 for row in cleaned_rows if row["date"] == day
            )
            writer.writerow({
                "date": day,
                "total_sales": f"{daily_sales[day]:.2f}",
                "valid_orders": order_count,
            })

    print("\n=== SALES CSV CLEANER: COMPLETE ===")
    print(f"Input records checked: {len(cleaned_rows) + len(rejected_rows)}")
    print(f"Valid records:         {len(cleaned_rows)}")
    print(f"Rejected records:      {len(rejected_rows)}")
    print(f"Total valid sales:     {sum(daily_sales.values()):.2f}")
    print("\nOutput files saved in:", OUTPUT_DIR)
    print("- cleaned_sales.csv")
    print("- rejected_rows.csv")
    print("- daily_sales_report.csv")


if __name__ == "__main__":
    try:
        process_sales()
    except OSError as error:
        print(f"ERROR: Could not create or write output files: {error}")
